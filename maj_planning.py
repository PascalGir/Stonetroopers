#!/usr/bin/env python3
"""Régénère data.js à partir du fichier Excel du planning.

Usage :  python3 maj_planning.py Planning26-27_Curling-Stonetroopes.xlsx
(nécessite : pip install openpyxl)
"""
import datetime as dt
import json
import sys

import openpyxl

PLAYERS = ["Suettania", "Adélie", "Océane", "Jean-Pierre", "Pascal"]


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def main(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Planning"]
    header = [c.value for c in ws[2]]
    cols = {name: header.index(name) for name in PLAYERS}
    matches = []
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        date = iso(r[0])
        if not date:
            continue
        heure = r[2].strftime("%H:%M") if isinstance(r[2], (dt.time, dt.datetime)) else str(r[2] or "")
        matches.append({
            "row": i,
            "date": date,
            "heure": heure,
            "piste": r[3] or "",
            "adversaire": r[4] or "",
            "dispo": {p: (str(r[cols[p]]).strip() if r[cols[p]] else "") for p in PLAYERS},
        })

    vacances = []
    if "Vacances" in wb.sheetnames:
        for r in wb["Vacances"].iter_rows(min_row=3, values_only=True):
            if r[0] and iso(r[1]) and iso(r[2]):
                vacances.append({"joueur": r[0], "debut": iso(r[1]), "fin": iso(r[2]), "commentaire": r[3] or ""})

    data = {"miseAJour": dt.date.today().isoformat(), "matches": matches, "vacances": vacances}
    with open("data.js", "w", encoding="utf-8") as f:
        f.write("// Généré par maj_planning.py — ne pas modifier à la main si tu utilises le script\n")
        f.write("window.PLANNING = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n")
    print(f"data.js écrit : {len(matches)} matchs, {len(vacances)} périodes de vacances")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Planning26-27_Curling-Stonetroopes.xlsx")
